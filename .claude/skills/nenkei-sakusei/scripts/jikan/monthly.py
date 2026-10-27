import openpyxl,datetime,re,copy
from openpyxl.styles import Font,PatternFill,Border
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
D=datetime.date; td=datetime.timedelta
WD='月火水木金土日'
wb=openpyxl.load_workbook('x/count.xlsm')          # 今年度の行事予定ファイル（書式の元）
master=openpyxl.load_workbook('/home/user/update/令和9年度_年間計画（案）.xlsx')
mws=master['年計']
HOL={D(2027,4,29):'昭和の日',D(2027,5,3):'憲法記念日',D(2027,5,4):'みどりの日',D(2027,5,5):'こどもの日',D(2027,7,19):'海の日',D(2027,8,11):'山の日',D(2027,9,20):'敬老の日',D(2027,9,23):'秋分の日',D(2027,10,11):'スポーツの日',D(2027,11,3):'文化の日',D(2027,11,23):'勤労感謝の日',D(2028,1,1):'元日',D(2028,1,10):'成人の日',D(2028,2,11):'建国記念の日',D(2028,2,23):'天皇誕生日',D(2028,3,20):'春分の日'}
def vac(d): return D(2027,7,21)<=d<=D(2027,8,31) or D(2027,12,24)<=d<=D(2028,1,6) or d>=D(2028,3,25) or d<D(2027,4,6)
def school(d): return d.weekday()<5 and d not in HOL and not vac(d) and d!=D(2027,11,26)
MCOL={4:'B',5:'E',6:'H',7:'K',8:'N',9:'Q',10:'T',11:'W',12:'Z',1:'AC',2:'AF',3:'AI'}
def nenkei(d):
    L=MCOL[d.month]; ec=gl(ci(L)+1); cc=gl(ci(L)+2); r=4+2*(d.day-1)
    ev=[];co=[]
    for rr in (r,r+1):
        c=mws[f'{ec}{rr}']
        if c.value:
            col=c.font.color.rgb if (c.font.color is not None and isinstance(c.font.color.rgb,str)) else None
            fill=c.fill.fgColor.rgb if c.fill.fill_type=='solid' and isinstance(c.fill.fgColor.rgb,str) and c.fill.fgColor.rgb not in ('FF92D050',) else None
            for x in str(c.value).split('・'):
                if x.strip(): ev.append((x.strip(),col,fill))
        v=mws[f'{cc}{rr}'].value
        if v: co+=[x for x in str(v).split('\n') if x]
    return ev,co
NAME={'職':'職員会議','学':'学年会','運':'運営委員会','現':'現職教育','進':'進路指導委員会','編':'学級編制会議','研':'研修'}
FIX={'定':'定時退校日','卒認':'卒業認定会議','進認':'進級認定会議','採':'教材採択委員会','全反':'年度末全体反省会','教支':'教育支援委員会','学保':'学校保健委員会','予':'予算委員会','不':'不登校生徒支援委員会','教':'教科部会','教・部活':'教科部会・部活','教･部活':'教科部会・部活','部会':'部会','努･採':'努力点推進委員会・教材採択委員会','(歓)':'歓迎会','▼評':'学校評議員会','旧運':'旧運営委員会','部反①':'年度末部会反省①','部反②':'年度末部会反省②','学6(2年)':'学年会6(2年)'}
def translate(d,co):
    other=[];student=[];shu=False;kenshu=False
    thu=d.weekday()==3
    for c in co:
        if c=='集': shu=True; continue
        if c in('委','議'):
            t='15:30' if (d.weekday() in(0,2) and '探½' in co) else '15:00'
            student.append(('生徒委員会' if c=='委' else '生徒議会')+f'({t})'); continue
        if c.startswith('探'):
            student.append('学習枠(13:50〜14:40)' if c!='探½' else '学習枠(14:45〜15:10)'); continue
        if c in FIX: other.append(FIX[c]); continue
        m=re.match(r'^(職|学|運|現|進|編|研)(\d+)$',c)
        if m:
            if m.group(1)=='研': other.insert(0,f'研修{m.group(2)}(14:55)'); kenshu=True
            else: other.append(f'{NAME[m.group(1)]}{m.group(2)}')
            continue
        other.append(c)
    # 時刻：木曜は研修の中、それ以外は15:00
    out=[]
    for o in other:
        if o.startswith('研修') or o=='定時退校日' or '(' in o: out.append(o)
        else: out.append(o+('' if (thu and kenshu) else '(15:00)'))
    return out,student,shu
CER=re.compile('始業式|終業式|修了式|入学式|卒業式|認証式|着任式')
# テンプレート：5月（31日分）の平日ブロックと休日ブロック
tpl=wb['R８5️⃣']
def snap(ws,r0):
    return {(dr,c):copy.copy(ws.cell(r0+dr,c)._style) for dr in range(3) for c in range(2,20)}
WEEK=snap(tpl,25); HOLI=snap(tpl,10)
months=[4,5,6,7,9,10,11,12,1,2,3]
EMO={4:'4️⃣',5:'5️⃣',6:'6️⃣',7:'7️⃣',9:'9️⃣',10:'🔟',11:'⑪',12:'⑫',1:'1️⃣',2:'2️⃣',3:'3️⃣'}
ZEN='０１２３４５６７８９'
import json as _j
PLAN=_j.load(open('plan.json')); SW2=_j.load(open('swap2.json')); POSD=_j.load(open('pos.json'))
newnames=[]
for m in months:
    y=2028 if m<=3 else 2027
    for student in (True,False):
        ws=wb.copy_worksheet(tpl)
        ws.title=f'R９{EMO[m]}'+(' (生徒用)' if student else '')
        newnames.append(ws.title)
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row>=7: ws.unmerge_cells(str(mr))
        for r in range(7,ws.max_row+1):
            for c in range(2,ci('AN')+1):
                cell=ws.cell(r,c); cell.value=None
                if c>=20: cell._style=copy.copy(ws.cell(200,c)._style)
        mm=f'{ZEN[m]}' if m<10 else str(m)
        ws['B3']=f'令和９年度　　     {mm} 月 行 事 予 定 表' if m<10 else f'令和９年度　　     {m}月 行 事 予 定 表'
        for day in range(1,32):
            top=7+3*(day-1)
            try: d=D(y,m,day)
            except ValueError:
                for dr in range(3):
                    for c in range(2,20): ws.cell(top+dr,c)._style=copy.copy(ws.cell(200,c)._style)
                    ws.row_dimensions[top+dr].height=4
                continue
            off=d.weekday()>=5 or d in HOL
            st=HOLI if off else WEEK
            for (dr,c),s in st.items(): ws.cell(top+dr,c)._style=copy.copy(s)
            for dr,h in enumerate((4.5,14.45,4.5) if off else (15.6,15.6,15.6)): ws.row_dimensions[top+dr].height=h
            ws.cell(top+1,2,day); ws.cell(top+1,3,WD[d.weekday()])
            ev,co=nenkei(d)
            other,stu,shu=translate(d,co)
            items=ev+[(s,None,None) for s in stu]
            # D：中段・下段に行事（上段は時間割変更用に空ける）
            lines=[i for i in items]
            if len(lines)>2: lines=[lines[0],('・'.join(x[0] for x in lines[1:]),lines[1][1],None)]
            for k,(t,col,fill) in enumerate(lines):
                c=ws.cell(top+1+k,4,t)
                c.font=Font(name=c.font.name,size=c.font.sz,color=col)
                if fill: c.fill=PatternFill('solid',fgColor=fill)
            if not student and other:
                fl=other if len(other)<=3 else other[:2]+['・'.join(other[2:])]
                for k,t in enumerate(fl): ws.cell(top+k,6,t)
            if school(d) and str(d) in PLAN:
                labs=PLAN[str(d)]; own=d.weekday()
                if any(not(x[0]==own and int(p)==x[1]) for p,x in labs.items()):
                    seq=[labs[str(p)] for p in sorted(map(int,labs))]; txt=''; prev=None
                    for x in seq:
                        txt+=(WD[x[0]] if x[0]!=prev else '')+ZEN[x[1]]; prev=x[0]
                    ws.cell(top,4,txt)
                    for gi,g in enumerate((1,2,3)):
                        pos=POSD.get(f'{d}|{g}') or []
                        for p in range(1,7):
                            if p in pos: ws.cell(top+gi,9+p,WD[labs[str(p)][0]]+ZEN[labs[str(p)][1]])
            for s_ in SW2:
                if s_['date']==str(d): ws.cell(top+1,9+s_['pos'],s_['mark'])
            if school(d):
                ws.cell(top+1,16,'Ｂ')
                s='・'.join(x[0] for x in ev)
                if 'ノー部活' in s: ws.cell(top+1,18,'×')
                elif d.weekday()!=3: ws.cell(top+1,18,'○')
                if CER.search(s): ws.cell(top+1,9,'式')
                elif shu: ws.cell(top+1,9,'◎')
        # 右側（印刷範囲外）に年計の抜き出し
        r=7
        for day in range(1,32):
            try: d=D(y,m,day)
            except ValueError: break
            ev,co=nenkei(d)
            ws.cell(r,22,day); ws.cell(r+1,22,WD[d.weekday()])
            ws.cell(r,23,'・'.join(x[0] for x in ev) or None); ws.cell(r,24,'・'.join(co) or None)
            r+=2
        ws.print_area='B3:R99'
# 古いシートを消し、年計をR9に差し替え
for n in list(wb.sheetnames):
    if n not in newnames and n!='年計': del wb[n]
nk=wb['年計']
for mr in list(nk.merged_cells.ranges): nk.unmerge_cells(str(mr))
for row in nk.iter_rows():
    for c in row: c.value=None
for row in mws.iter_rows():
    for c in row:
        t=nk.cell(c.row,c.column); t.value=c.value
        if c.has_style: t._style=copy.copy(c._style)
for mr in mws.merged_cells.ranges: nk.merge_cells(str(mr))
for k,v in mws.column_dimensions.items(): nk.column_dimensions[k].width=v.width
for v in wb.views: v.firstSheet=0; v.activeTab=0
for nm in list(wb.defined_names.keys()): del wb.defined_names[nm]
for wsx in wb.worksheets:
    for nm in list(wsx.defined_names.keys()):
        if nm!='_xlnm.Print_Area': del wsx.defined_names[nm]
    for row in wsx.iter_rows():
        for c in row:
            if isinstance(c.value,str) and ('WEBSERVICE' in c.value or c.value.startswith('ERROR')): c.value=None
out='/home/user/update/令和9年度_月行事予定（案）.xlsx'
wb.save(out); print(out, wb.sheetnames)
