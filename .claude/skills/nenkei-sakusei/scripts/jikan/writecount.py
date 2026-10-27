import openpyxl,json,datetime,re
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as gl, column_index_from_string as ci
D=datetime.date; td=datetime.timedelta
WD='月火水木金土日'; Z='０１２３４５６'
P=json.load(open('pos.json')); L=json.load(open('plan.json'))
master=openpyxl.load_workbook('/home/user/update/令和9年度_年間計画（案）.xlsx')['年計']
MC={4:'B',5:'E',6:'H',7:'K',8:'N',9:'Q',10:'T',11:'W',12:'Z',1:'AC',2:'AF',3:'AI'}
def evs(d):
    ec=gl(ci(MC[d.month])+1); r=4+2*(d.day-1)
    return '・'.join(str(master[f'{ec}{rr}'].value) for rr in (r,r+1) if master[f'{ec}{rr}'].value)
def lab(x): return WD[x[0]]+Z[x[1]]
F=Font(name='メイリオ',size=9); FB=Font(name='メイリオ',size=9,bold=True)
th=Side(style='thin',color='999999'); B=Border(left=th,right=th,top=th,bottom=th)
hl=PatternFill('solid',fgColor='FFF2CC'); hd=PatternFill('solid',fgColor='DDE6F0'); gr=PatternFill('solid',fgColor='D9D9D9')
wb=openpyxl.Workbook(); ws=wb.active; ws.title='授業数カウント'
ws.append(['令和9年度 授業数カウント（B帯・45分）　黄色＝その曜日の時間割と違うコマ　灰色＝授業なし'])
ws.append(['日付','曜','行事（年計）','時間割変更']+[f'{g}年{p}限' for g in (1,2,3) for p in range(1,7)])
d=D(2027,4,1); r=3; first=3
while d<=D(2028,3,31):
    ds=str(d)
    if ds in L:
        ws.cell(r,1,d).number_format='m/d'; ws.cell(r,2,WD[d.weekday()]); ws.cell(r,3,evs(d))
        labs=L[ds]; own=d.weekday()
        changed=any(not(x[0]==own and int(p)==x[1]) for p,x in labs.items())
        if changed:
            # 例：金１２３４水６金６
            seq=[labs[str(p)] for p in sorted(map(int,labs))]
            txt=''; prev=None
            for x in seq:
                txt+=(WD[x[0]] if x[0]!=prev else '')+Z[x[1]]; prev=x[0]
            ws.cell(r,4,txt)
        for gi,g in enumerate((1,2,3)):
            pos=P.get(f'{ds}|{g}') or []
            for p in range(1,7):
                c=ws.cell(r,5+gi*6+p-1)
                if p in pos:
                    x=labs[str(p)]; c.value=lab(x)
                    if not(x[0]==own and x[1]==p): c.fill=hl
                elif P.get(f'{ds}|{g}') is not None and not(d.weekday()==3 and p==6): c.fill=gr
        r+=1
    d+=td(days=1)
last=r-1
for row in ws.iter_rows(min_row=2,max_row=last):
    for c in row: c.font=FB if c.row==2 else F; c.border=B; c.alignment=Alignment(horizontal='center',vertical='center',shrink_to_fit=True)
for c in ws[2]: c.fill=hd
ws.column_dimensions['C'].width=34; ws.column_dimensions['D'].width=16
for i in range(5,23): ws.column_dimensions[gl(i)].width=5.2
ws.freeze_panes='E3'
# 区切り別集計（数式）
ss=wb.create_sheet('区切り別集計')
ENDS=[(D(2027,4,6),D(2027,5,19),'4/6〜確認テスト'),(D(2027,5,20),D(2027,6,28),'確認テスト後〜1学期期末'),(D(2027,6,29),D(2027,9,28),'1学期期末後〜中間'),(D(2027,9,29),D(2027,11,18),'中間後〜2学期期末'),(D(2027,11,19),D(2028,2,18),'2学期期末後〜学年末'),(D(2028,2,19),D(2028,3,24),'学年末後〜修了')]
SL=[(w,p) for w in range(5) for p in range(1,7) if not(w==3 and p==6)]
ss.append(['各コマの実施回数（区切りごと）。差は2以内が目標'])
row=3
for g in (1,2,3):
    gc1=gl(5+(g-1)*6); gc2=gl(10+(g-1)*6)
    ss.cell(row,1,f'{g}年').font=FB; row+=1
    ss.cell(row,1,'区切り')
    for j,s in enumerate(SL): ss.cell(row,2+j,lab(s))
    ss.cell(row,2+len(SL),'最少'); ss.cell(row,3+len(SL),'最多'); ss.cell(row,4+len(SL),'差')
    for c in ss[row]: c.fill=hd; c.font=FB
    hdr=row; row+=1
    for a,b,name in ENDS+[(D(2027,4,6),D(2028,3,24),'年間')]:
        ss.cell(row,1,name)
        for j,s in enumerate(SL):
            ss.cell(row,2+j,f'=SUMPRODUCT((授業数カウント!$A$3:$A${last}>=DATE({a.year},{a.month},{a.day}))*(授業数カウント!$A$3:$A${last}<=DATE({b.year},{b.month},{b.day}))*(授業数カウント!${gc1}$3:${gc2}${last}="{lab(s)}"))')
        e=gl(1+len(SL))
        ss.cell(row,2+len(SL),f'=MIN(B{row}:{e}{row})'); ss.cell(row,3+len(SL),f'=MAX(B{row}:{e}{row})'); ss.cell(row,4+len(SL),f'={gl(3+len(SL))}{row}-{gl(2+len(SL))}{row}')
        row+=1
    row+=1
for r_ in ss.iter_rows(min_row=3,max_row=row):
    for c in r_:
        if c.value is not None: c.border=B; c.alignment=Alignment(horizontal='center'); c.font=FB if c.font.bold else F
ss.column_dimensions['A'].width=22
for i in range(2,34): ss.column_dimensions[gl(i)].width=5
mm=wb.create_sheet('作り方')
for t in ['令和9年度 授業数カウントの作り方',
 '・年計の行事から、毎日・学年ごとに授業ができる時限を出し、どの曜日の何限の時間割を行うかを割り振った。',
 '・条件1：確認テスト・1学期期末・中間・2学期期末・学年末テストの区切りごとに、各コマの回数の差を2以内にする。',
 '・条件2：1限には、どこかの曜日の1限か、6限の道徳・学活・総合（月6・水6・金6）だけを置く（2限を1限に繰り上げない）。',
 '・条件3：1日に使う曜日は2つまで。',
 '・各時限は、全学年同じ曜日・時限を行う（先生の重なりを防ぐため）。',
 '・結果：3年の「2学期期末後〜学年末」だけ差3。ほかはすべて差2以内。']:
    mm.append([t])
mm.column_dimensions['A'].width=110
out='/home/user/update/令和9年度_授業数カウント（案）.xlsx'
wb.save(out); print(out,last)
