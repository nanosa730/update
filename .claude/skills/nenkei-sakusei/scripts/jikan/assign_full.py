import openpyxl,re,datetime,json,collections,itertools
D=datetime.date; td=datetime.timedelta
WD='月火水木金'
# ---------- 1. 今年度（R8）の各日・学年の授業時限 ----------
wbv=openpyxl.load_workbook('x/gyoji.xlsx',data_only=True)['年計']
band=json.load(open('band.json'))
slotre=re.compile(r'^[月火水木金][１２３４５６1-6]$')
G={1:range(4,10),2:range(10,16),3:range(16,22)}
R8P={}
for r in range(3,wbv.max_row+1):
    d=wbv.cell(r,1).value
    if not isinstance(d,(int,float)): continue
    d=D(1899,12,30)+td(days=int(d))
    if d.weekday()>=5 or band.get(str(d),'')=='': continue
    for g,cols in G.items():
        v=[re.sub(r'\s|　','',str(wbv.cell(r,c).value or '')) for c in cols]
        P=[i+1 for i,x in enumerate(v) if slotre.match(x)]
        if d.weekday()!=3 and P==[1,2,3,4,5] and v[5]=='': P=[1,2,3,4,5,6]
        R8P[(d,g)]=(P, d.weekday()==3)
# ---------- 2. 令和9年度の各日・学年の授業時限 ----------
HOL={D(2027,4,29),D(2027,5,3),D(2027,5,4),D(2027,5,5),D(2027,7,19),D(2027,9,20),D(2027,9,23),D(2027,10,11),D(2027,11,3),D(2027,11,23),D(2028,1,10),D(2028,2,11),D(2028,2,23),D(2028,3,20),D(2027,11,26)}
def vac(d): return D(2027,7,21)<=d<=D(2027,8,31) or D(2027,12,24)<=d<=D(2028,1,6) or d>=D(2028,3,25) or d<D(2027,4,6)
def school(d): return d.weekday()<5 and d not in HOL and not vac(d)
src={D(2027,4,6):D(2026,4,6),D(2027,4,7):D(2026,4,7),D(2027,4,8):D(2026,4,8),D(2027,4,9):D(2026,4,9),D(2027,4,12):D(2026,4,10),
     D(2027,7,20):D(2026,7,17),D(2027,7,16):D(2026,7,16),D(2027,7,15):D(2026,7,15),D(2027,9,1):D(2026,9,1),D(2027,12,23):D(2026,12,23),D(2027,12,22):D(2026,12,22),D(2027,12,21):D(2026,12,21),
     D(2028,1,7):D(2027,1,7),D(2028,3,24):D(2027,3,24),D(2028,3,23):D(2027,3,23),
     D(2027,5,10):D(2026,5,8),D(2027,5,11):D(2026,5,11),D(2027,5,7):D(2026,5,7),
     D(2027,11,2):D(2026,11,4),D(2027,11,17):D(2026,11,17),D(2027,11,18):D(2026,11,18),
     D(2028,2,17):D(2027,2,17),D(2028,2,18):D(2027,2,18)}
NORMAL={D(2027,4,22),D(2027,5,6),D(2027,6,4),D(2027,7,9),D(2027,11,15),D(2027,12,7),D(2027,11,16),D(2028,2,16),D(2028,3,21),D(2028,3,22)}
FULL='full'
over={}
for d in (D(2027,5,25),D(2027,5,26),D(2027,5,27),D(2027,5,28)): over[(d,2)]=FULL
for d in (D(2027,6,14),D(2027,6,15),D(2027,6,16)): over[(d,3)]=FULL
for d in (D(2027,6,8),D(2027,6,9),D(2027,6,10),D(2027,6,11)): over[(d,3)]=[]
over[(D(2027,6,7),3)]=R8P[(D(2026,6,12),3)][0]
over[(D(2027,4,15),3)]=[4,5]
for g in (2,3): over[(D(2027,4,7),g)]=[1,2,3,4]
over[(D(2027,5,7),3)]=[1,2,3,4]
over[(D(2027,4,30),3)]=FULL
over[(D(2027,11,17),3)]=[]; over[(D(2027,11,18),3)]=[]
for d in (D(2028,2,17),D(2028,2,18)): over[(d,3)]=FULL
over[(D(2028,3,1),3)]=FULL; over[(D(2028,3,2),3)]=FULL
over[(D(2028,3,3),3)]=R8P[(D(2027,3,3),3)][0]; over[(D(2028,3,6),3)]=R8P[(D(2027,3,4),3)][0]; over[(D(2028,3,7),3)]=R8P.get((D(2027,3,5),3),([],0))[0]
def positions(d,g):
    cap=5 if d.weekday()==3 else 6
    full=list(range(1,cap+1))
    if g==3 and d>D(2028,3,7): return None

    if (d,g) in over:
        v=over[(d,g)]; P=full if v==FULL else list(v)
    elif d in NORMAL: P=full
    else:
        s=src.get(d,d-td(days=364))
        if (s,g) not in R8P: P=full
        else:
            P,wasthu=R8P[(s,g)]
            if P==list(range(1,(5 if wasthu else 6)+1)): P=full
            P=[p for p in P if p<=cap]
    if d in (D(2027,7,16),D(2027,12,22),D(2028,3,23)) and P:   # 大掃除の日は1コマ増
        nxt=max(P)+1
        if nxt<=cap: P=P+[nxt]
    return P
days=[]
d=D(2027,4,1)
while d<=D(2028,3,31):
    if school(d): days.append(d)
    d+=td(days=1)
POS={(d,g):positions(d,g) for d in days for g in G}
tot={g:sum(len(POS[(d,g)] or []) for d in days) for g in G}
print('授業コマ',tot)
json.dump({f'{d}|{g}':POS[(d,g)] for d in days for g in G},open('pos_full.json','w'))
