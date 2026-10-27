import openpyxl,re,sys
from openpyxl.styles import Alignment,Font
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
p=sys.argv[1]
wb=openpyxl.load_workbook(p); ws=wb['年計']
CODE=re.compile(r'^(研\d+|職\d+|学\d+|学6\(2年\)|運\d+|旧運|現\d+|進\d+|編\d|委|議|集|定|教|予|不|採|卒認|進認|全反|部反.|学保|教支|部会|教[・･]部活|努･採|\(歓\)|▼評)$')
moved=[]
for L in ['B','E','H','K','N','Q','T','W','Z','AC','AF','AI']:
    ec,cc=gl(ci(L)+1),gl(ci(L)+2)
    ws.column_dimensions[cc].width=5.5
    for day in range(1,32):
        r1=4+2*(day-1)
        notes=[]
        for r in (r1,r1+1):
            v=ws[f'{cc}{r}'].value
            if not isinstance(v,str): continue
            keep=[]
            for line in v.split('\n'):
                s=line.strip(' 　')
                if not s: continue
                (keep if CODE.match(s) else notes).append(s)
            ws[f'{cc}{r}'].value='\n'.join(keep) or None
        for n in notes:
            for r in (r1,r1+1):
                c=ws[f'{ec}{r}']
                if not c.value: c.value=n; break
            else:
                c=ws[f'{ec}{r1+1}']; c.value=c.value+'・'+n
            moved.append((L,day,n))
        for r in (r1,r1+1):
            c=ws[f'{cc}{r}']; multi=isinstance(c.value,str) and '\n' in c.value
            c.alignment=Alignment(horizontal='right',vertical='center',wrap_text=multi,shrink_to_fit=not multi)
            c.font=Font(name=c.font.name,size=7 if multi else 9,color=c.font.color)
wb.save(p); print(len(moved)); [print(m) for m in moved]
